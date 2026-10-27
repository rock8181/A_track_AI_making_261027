"""공정/설비 데이터에서 이상데이터를 찾아 엑셀 보고서로 정리

이상 기준 (하나라도 해당하면 이상):
  1) 불량률 >= RATE_LIMIT
  2) 평균값이 정상범위(하한~상한)를 벗어남
  3) 최대값/최소값이 정상범위를 (상한-하한) × EXCURSION_RATIO 넘게 벗어남

사용법:
  python anomaly_report.py                       # 아래 CONFIG 그대로 실행
  python anomaly_report.py --inputs a.csv b.csv --key LOT_ID --out 결과.xlsx
  python anomaly_report.py --inputs a.csv b.csv --inspect   # 컬럼·인코딩만 확인
필요 패키지: pandas, openpyxl
"""
import argparse
import os
import re
import sys

# ── 기준값 (여기만 바꿔서 다시 실행) ──
RATE_LIMIT = 3.0        # 1) 불량률(%) 기준
EXCURSION_RATIO = 0.6   # 3) 극값 허용폭 = (상한-하한) × 이 값

# ── 입력/컬럼 설정 (스크립트를 결과 폴더에 복사한 뒤 이 블록을 맞춰 두면 인자 없이 재실행 가능) ──
CONFIG = {
    "inputs": ["data/검사로그.csv", "data/설비센서로그.csv"],  # 1개 또는 여러 개 (csv/xlsx/xls)
    "key": "LOT_ID",              # 여러 파일을 붙일 키 (파일 1개면 무시)
    "out": "이상데이터_분석.xlsx",
    "rate_col": "불량률(%)",      # 없으면 기준 1) 생략
    "mean_col": "평균값",         # 없으면 기준 2) 생략
    "max_col": "최대값",
    "min_col": "최소값",          # 최대/최소 둘 다 없으면 기준 3) 생략
    "range_col": "정상범위",      # "하한~상한" 형식 한 컬럼
    "lower_col": None,            # 또는 하한/상한이 따로 있으면 컬럼명 지정 (range_col보다 우선)
    "upper_col": None,
    "date_col": "검사일자",       # 월별 추이·정렬용 (없으면 생략)
    "group_cols": [["설비ID", "공정", "라인"], ["공정", "모니터링항목", "정상범위"]],  # 집계 시트 (없는 컬럼은 자동 제외)
}

ENCODINGS = ["utf-8-sig", "cp949", "euc-kr", "latin-1"]

# Windows 콘솔(cp949) 등에서 한글 출력이 깨지거나 죽지 않게
for s in (sys.stdout, sys.stderr):
    try:
        s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

try:
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as e:
    sys.exit(f"[오류] 필요한 패키지가 없습니다: {e.name}\n  설치: python -m pip install pandas openpyxl")


def read_table(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(path), "excel"
    for enc in ENCODINGS:
        try:
            return pd.read_csv(path, encoding=enc, sep=None, engine="python"), enc
        except UnicodeDecodeError:
            continue
    sys.exit(f"[오류] 인코딩을 판별할 수 없습니다: {path}")


NUM = r"[-+]?\d+(?:\.\d+)?"
RANGE_RE = re.compile(rf"^\s*({NUM})\s*(?:~|∼|〜|-|–|—|,|/)\s*({NUM})\s*$")


def parse_range(v):
    m = RANGE_RE.match(str(v))
    if not m:
        return (float("nan"), float("nan"))
    lo, hi = float(m.group(1)), float(m.group(2))
    return (min(lo, hi), max(lo, hi))


def main():
    ap = argparse.ArgumentParser(description="이상데이터 엑셀 보고서")
    ap.add_argument("--inputs", nargs="+")
    ap.add_argument("--out")
    ap.add_argument("--key")
    for k in ("rate_col", "mean_col", "max_col", "min_col", "range_col", "lower_col", "upper_col", "date_col"):
        ap.add_argument("--" + k.replace("_", "-"), dest=k)
    ap.add_argument("--rate-limit", type=float, default=RATE_LIMIT)
    ap.add_argument("--excursion", type=float, default=EXCURSION_RATIO)
    ap.add_argument("--inspect", action="store_true", help="파일별 컬럼·인코딩·샘플만 출력")
    a = ap.parse_args()
    cfg = {**CONFIG, **{k: v for k, v in vars(a).items() if v is not None and k in CONFIG}}

    # 상대경로는 스크립트가 아니라 현재 작업 폴더 기준
    frames = []
    for p in cfg["inputs"]:
        if not os.path.exists(p):
            sys.exit(f"[오류] 입력 파일이 없습니다: {os.path.abspath(p)}")
        df_, enc = read_table(p)
        df_.columns = [str(c).strip() for c in df_.columns]
        frames.append(df_)
        if a.inspect:
            print(f"== {p}  ({len(df_)}행, 인코딩 {enc})\n컬럼: {list(df_.columns)}\n{df_.head(3).to_string()}\n")
    if a.inspect:
        return

    key = cfg["key"]
    df = frames[0]
    for other in frames[1:]:
        if key not in df.columns or key not in other.columns:
            sys.exit(f"[오류] 결합 키 '{key}'가 모든 파일에 있어야 합니다. --key로 지정하세요.")
        dup = [c for c in other.columns if c in df.columns and c != key]
        before = len(df)
        df = df.merge(other.drop(columns=dup), on=key, how="outer", indicator=True)
        miss = (df["_merge"] != "both").sum()
        if miss:
            print(f"[주의] 키가 한쪽에만 있는 행 {miss}건 (결합 전 {before}행)")
        df = df.drop(columns="_merge")
    if len(frames) > 1 and df[key].duplicated().any():
        print(f"[주의] 키 '{key}'가 중복된 행이 있습니다 (1:N 결합). 결과 행 수: {len(df)}")
    df = df.reset_index(drop=True)

    has = lambda c: c is not None and c in df.columns
    num = lambda c: pd.to_numeric(df[c], errors="coerce")
    false = pd.Series(False, index=df.index)

    # 정상범위
    if has(cfg["lower_col"]) and has(cfg["upper_col"]):
        lo, hi = num(cfg["lower_col"]), num(cfg["upper_col"])
        range_label = df[cfg["lower_col"]].astype(str) + "~" + df[cfg["upper_col"]].astype(str)
    elif has(cfg["range_col"]):
        r = df[cfg["range_col"]].map(parse_range)
        lo, hi = r.str[0], r.str[1]
        range_label = df[cfg["range_col"]].astype(str)
        bad = lo.isna().sum()
        if bad:
            print(f"[주의] 정상범위 형식을 읽지 못한 행 {bad}건 → 기준 2)·3) 판정 제외. 예: {df.loc[lo.isna(), cfg['range_col']].unique()[:5].tolist()}")
    else:
        lo = hi = None
    margin = (hi - lo) * a.excursion if lo is not None else None

    skipped = []
    c1 = num(cfg["rate_col"]) >= a.rate_limit if has(cfg["rate_col"]) else (skipped.append("1) 불량률") or false)
    if lo is not None and has(cfg["mean_col"]):
        m = num(cfg["mean_col"])
        c2 = (m < lo) | (m > hi)
    else:
        skipped.append("2) 평균값"); c2 = false
    c3_max = num(cfg["max_col"]) > hi + margin if lo is not None and has(cfg["max_col"]) else false
    c3_min = num(cfg["min_col"]) < lo - margin if lo is not None and has(cfg["min_col"]) else false
    if lo is None or not (has(cfg["max_col"]) or has(cfg["min_col"])):
        skipped.append("3) 최대/최소값")
    c3 = c3_max | c3_min
    if skipped:
        print(f"[주의] 컬럼이 없어 생략한 기준: {', '.join(skipped)}")

    def reasons(i):
        r = []
        if c1[i]: r.append(f"불량률 {df.at[i, cfg['rate_col']]} ≥ {a.rate_limit}")
        if c2[i]: r.append(f"평균값 {df.at[i, cfg['mean_col']]} 정상범위({range_label[i]}) 이탈")
        if c3_max[i]: r.append(f"최대값 {df.at[i, cfg['max_col']]} > 상한+{margin[i]:.4g}")
        if c3_min[i]: r.append(f"최소값 {df.at[i, cfg['min_col']]} < 하한-{margin[i]:.4g}")
        return " / ".join(r)

    flags = {"①불량률": c1, "②평균이탈": c2, "③극값이탈": c3}
    for name, s in flags.items():
        df[name] = s.map({True: "●", False: ""})
    df["해당기준수"] = c1.astype(int) + c2.astype(int) + c3.astype(int)
    is_anom = df["해당기준수"] > 0
    anom = df[is_anom].copy()
    anom["이상사유"] = [reasons(i) for i in anom.index]
    date_col = cfg["date_col"] if has(cfg["date_col"]) else None
    if date_col:
        anom = anom.sort_values([date_col] + ([key] if key in anom.columns else []))

    n = len(df)
    summary = pd.DataFrame([
        [f"① 불량률 ≥ {a.rate_limit}", int(c1.sum())],
        ["② 평균값이 정상범위(하한~상한) 벗어남", int(c2.sum())],
        [f"③ 최대/최소값이 정상범위를 (상한-하한)×{a.excursion} 넘게 벗어남", int(c3.sum())],
        ["이상 행 합계 (하나라도 해당, 중복 제외)", int(is_anom.sum())],
        ["전체 행 수", n],
        ["이상 비율(%)", round(is_anom.sum() / n * 100, 2) if n else 0],
    ], columns=["구분", "건수"], dtype=object)  # 건수가 108.0처럼 실수로 바뀌지 않게

    def group(cols):
        g = df.groupby(cols, dropna=False).agg(
            전체=("해당기준수", "size"),
            불량률이상=("①불량률", lambda s: (s == "●").sum()),
            평균이탈=("②평균이탈", lambda s: (s == "●").sum()),
            극값이탈=("③극값이탈", lambda s: (s == "●").sum()),
            이상건수=("해당기준수", lambda s: (s > 0).sum()))
        g["이상비율(%)"] = (g["이상건수"] / g["전체"] * 100).round(2)
        return g.reset_index()

    sheets = {"요약": summary, "이상데이터": anom}
    for cols in cfg["group_cols"]:
        if not cols or cols[0] not in df.columns:  # 대표 컬럼이 없으면 그 집계는 생략
            continue
        cols = [c for c in cols if c in df.columns]
        name = "_".join(cols[:2]) + "별"
        sheets[re.sub(r"[\[\]:*?/\\]", "", name)[:31]] = group(cols).sort_values("이상건수", ascending=False)
    if date_col:
        dt = pd.to_datetime(df[date_col], errors="coerce")
        if dt.notna().any():
            df["연월"] = dt.dt.strftime("%Y-%m")
            sheets["월별추이"] = group("연월").sort_values("연월")

    out = cfg["out"]
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    try:
        with pd.ExcelWriter(out, engine="openpyxl") as w:
            for name, d in sheets.items():
                d.to_excel(w, sheet_name=name, index=False)
    except PermissionError:
        sys.exit(f"[오류] 파일에 쓸 수 없습니다(엑셀에서 열려 있는지 확인): {os.path.abspath(out)}")

    # 서식
    wb = load_workbook(out)
    hdr = PatternFill("solid", fgColor="1F4E78")
    red = PatternFill("solid", fgColor="F8CBAD")
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = hdr
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, 1):
            w_ = max(len(str(c.value or "")) for c in list(col)[:300])
            ws.column_dimensions[get_column_letter(i)].width = min(max(w_ * 1.6 + 2, 8), 70)
        if ws.title != "요약" and ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

    ws = wb["이상데이터"]
    h = {c.value: c.column for c in ws[1]}
    for r, i in enumerate(anom.index, start=2):
        for col, s in ((cfg["rate_col"], c1), (cfg["mean_col"], c2), (cfg["max_col"], c3_max), (cfg["min_col"], c3_min)):
            if s[i] and col in h:
                ws.cell(r, h[col]).fill = red
        for k in (*flags, "해당기준수"):
            ws.cell(r, h[k]).alignment = Alignment(horizontal="center")

    ws = wb["요약"]
    ws.column_dimensions["A"].width = 60
    notes = [f"기준값: 불량률 ≥ {a.rate_limit}, 극값 허용폭 = (상한-하한) × {a.excursion}",
             f"원천: {', '.join(cfg['inputs'])}" + (f" ({key} 기준 결합)" if len(frames) > 1 else "")]
    if skipped:
        notes.append(f"생략한 기준(컬럼 없음): {', '.join(skipped)}")
    for j, t in enumerate(notes):
        ws.cell(len(summary) + 3 + j, 1, t).font = Font(italic=True, color="666666")
    wb.save(out)

    print(summary.to_string(index=False))
    print(f"\n저장: {os.path.abspath(out)}  (시트: {', '.join(sheets)})")


if __name__ == "__main__":
    main()
