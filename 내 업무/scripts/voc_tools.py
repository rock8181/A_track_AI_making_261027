"""VOC 처리 Workflow의 코드 Task 도구 (workflow.md의 코드 1 · 3 · 6).

`내 업무/` 폴더에서 실행한다.

  python scripts/voc_tools.py load      # 코드 1: 티켓 + 요청자 정보 조회 → output/work/tickets.json
                                        #         (과거 처리 이력도 AI가 읽을 수 있게 history.json으로 변환)
  python scripts/voc_tools.py context   # 코드 3: 티켓별 로그·DB 조회 → output/work/context/<티켓>.json
                                        #         (AI 2 결과 output/work/classified.json 필요)
  python scripts/voc_tools.py status    # 티켓별 진행 상황 (Router 경로, 분석·답변·검토 상태)
  python scripts/voc_tools.py report    # 코드 6: 결과 취합 → output/VOC_검토보고서_<YYYYMMDD>.html

AI Task(2 · 4 · 5 · 5-1)는 .claude/agents/의 서브에이전트가 output/work/ 아래에 JSON으로 결과를 남긴다.
"""

import csv
import html
import json
import re
import shutil
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
TICKET_FILE = DATA_DIR / "VOC_접수티켓.xlsx"
HISTORY_FILE = DATA_DIR / "과거VOC처리이력.xlsx"
DB_DIR = DATA_DIR / "시스템DB"
OUTPUT_DIR = BASE_DIR / "output"
WORK_DIR = OUTPUT_DIR / "work"

OPEN_STATUS = "접수"          # 처리 대상 티켓 상태
CONTEXT_DAYS = 3              # 접수 시각 기준 며칠 전 로그까지 조회할지
MAX_REVISIONS = 2             # Gen / Eval 재작성 최대 횟수

# 처리 기한(SLA) — 운영 매뉴얼 9장
SLA_HIGH_HOURS = 2            # 높음: 접수 후 2시간
SLA_DAY_END = "18:00"         # 보통: 당일 업무 종료 시각
SLA_LOW_BUSINESS_DAYS = 3     # 낮음: 3영업일 이내
SOON_MINUTES = 30             # 기한까지 이 시간 이하로 남으면 '임박'

ROUTES = ("장애·문의", "중복 VOC", "개선 요청")
PR_ID_PATTERN = re.compile(r"PR-\d{8}")


# ---------------------------------------------------------------- 공통

def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def read_csv(name):
    with (DB_DIR / name).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_xlsx(path):
    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    return [dict(zip(header, ["" if v is None else str(v) for v in row])) for row in rows[1:] if any(row)]


def parse_dt(text):
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(text.strip(), fmt)
        except ValueError:
            continue
    raise ValueError(f"날짜 형식을 알 수 없음: {text}")


def load_classified():
    path = WORK_DIR / "classified.json"
    if not path.exists():
        sys.exit("output/work/classified.json이 없습니다. voc-classifier(AI 2)를 먼저 실행하세요.")
    return {c["ticket_id"]: c for c in read_json(path)["tickets"]}


# ---------------------------------------------------------------- 코드 1: 티켓 + 요청자 정보

def cmd_load():
    users = {u["user_id"]: u for u in read_csv("users.csv")}
    tickets = []
    for row in read_xlsx(TICKET_FILE):
        if row["상태"] != OPEN_STATUS:
            continue
        tickets.append({
            "ticket_id": row["티켓번호"],
            "received_at": row["접수일시"],
            "requester": users.get(row["요청자ID"], {"user_id": row["요청자ID"], "name": row["요청자"],
                                                       "dept": row["부서"], "status": "사용자 DB에 없음"}),
            "category": row["카테고리(요청자 선택)"],
            "title": row["제목"],
            "content": row["내용"],
            "priority": row["우선순위"],
        })

    # 이전 실행 결과가 섞이지 않도록 작업 폴더를 새로 만든다
    if WORK_DIR.exists():
        shutil.rmtree(WORK_DIR)
    write_json(WORK_DIR / "tickets.json", {"loaded_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
                                           "tickets": tickets})
    write_json(WORK_DIR / "history.json", read_xlsx(HISTORY_FILE))
    print(f"티켓 {len(tickets)}건 → output/work/tickets.json")
    print("과거 처리 이력 → output/work/history.json")


# ---------------------------------------------------------------- 코드 3: 티켓별 로그·DB 조회

def cmd_context():
    tickets = read_json(WORK_DIR / "tickets.json")["tickets"]
    classified = load_classified()
    users = {u["user_id"]: u for u in read_csv("users.csv")}
    prs = read_csv("purchase_request.csv")
    logs = read_csv("error_log.csv")
    codes = read_csv("code_master.csv")

    for t in tickets:
        tid = t["ticket_id"]
        cls = classified.get(tid)
        if cls is None:
            sys.exit(f"classified.json에 {tid}가 없습니다.")
        uid = t["requester"]["user_id"]
        received = parse_dt(t["received_at"])
        since = received - timedelta(days=CONTEXT_DAYS)

        # 대상 사용자: 요청자 + AI 2가 찾은 실제 대상자(팀장이 팀원 대신 문의한 경우 등)
        target_ids = {uid, *cls.get("target_user_ids", [])}
        req_ids = set(cls.get("mentioned_req_ids", [])) | set(PR_ID_PATTERN.findall(t["content"]))

        related_prs = [p for p in prs if p["req_id"] in req_ids
                       or p["requester_id"] in target_ids or p["current_approver_id"] in target_ids]
        approver_ids = {p["current_approver_id"] for p in related_prs if p["current_approver_id"]}
        inactive_approvers = [a for a in approver_ids if users.get(a, {}).get("status") == "INACTIVE"]
        same_approver_pending = [p for p in prs if p["current_approver_id"] in inactive_approvers
                                 and p["status"] == "PENDING_APPROVAL"]
        # 결재선 재지정 후보: 요청자 부서의 사용 중인 결재자
        dept_approvers = [u for u in users.values() if u["dept"] == t["requester"].get("dept")
                          and u["role"] == "APPROVER" and u["status"] == "ACTIVE"]

        pr_ids = {p["req_id"] for p in related_prs}
        in_window = [l for l in logs if since <= parse_dt(l["occurred_at"]) <= received]
        user_logs = [l for l in in_window if l["user_id"] in target_ids or l["req_id"] in pr_ids]
        batch_logs = [l for l in in_window if l["user_id"] == "SYSTEM"]

        item_codes = {p["item_code"] for p in related_prs}
        code_rows = [c for c in codes if c["code_type"] != "ITEM"
                     or c["code"] in item_codes or c["use_yn"] == "N"]

        write_json(WORK_DIR / "context" / f"{tid}.json", {
            "ticket": t,
            "classification": cls,
            "target_users": [users[i] for i in sorted(target_ids) if i in users],
            "purchase_requests": related_prs,
            "approvers": [users[a] for a in sorted(approver_ids) if a in users],
            "same_approver_pending": same_approver_pending,
            "dept_active_approvers": dept_approvers,
            "error_logs": user_logs,
            "batch_logs": batch_logs,
            "code_master": code_rows,
            "query_window": f"{since:%Y-%m-%d %H:%M} ~ {received:%Y-%m-%d %H:%M}",
        })
    print(f"티켓 {len(tickets)}건 조회 결과 → output/work/context/")


# ---------------------------------------------------------------- 진행 상황

def ticket_state(tid, cls):
    """티켓 1건의 Gen / Eval 진행 상태를 계산한다."""
    analysis = WORK_DIR / "analysis" / f"{tid}.json"
    answer_path = WORK_DIR / "answer" / f"{tid}.json"
    review_path = WORK_DIR / "review" / f"{tid}.json"
    answer = read_json(answer_path) if answer_path.exists() else None
    review = read_json(review_path) if review_path.exists() else None

    if cls["route"] == "장애·문의" and not analysis.exists():
        return "분석 필요", answer, review
    if answer is None:
        return "답변 작성 필요", answer, review
    if review is None or review.get("revision") != answer.get("revision"):
        return "검토 필요", answer, review
    if review["result"] == "PASS":
        return "검토 통과", answer, review
    if answer.get("revision", 0) >= MAX_REVISIONS:
        return "검토 미통과", answer, review
    return "재작성 필요", answer, review


def cmd_status():
    classified = load_classified()
    for tid, cls in classified.items():
        state, answer, _ = ticket_state(tid, cls)
        rev = "-" if answer is None else answer.get("revision", 0)
        print(f"{tid}\t{cls['route']}\t재작성 {rev}회\t{state}")


# ---------------------------------------------------------------- 코드 6: 검토 보고서

def deadline_of(received, priority):
    if priority == "높음":
        return received + timedelta(hours=SLA_HIGH_HOURS)
    end_h, end_m = map(int, SLA_DAY_END.split(":"))
    if priority == "보통":
        return received.replace(hour=end_h, minute=end_m, second=0)
    day, left = received, SLA_LOW_BUSINESS_DAYS
    while left:
        day += timedelta(days=1)
        if day.weekday() < 5:
            left -= 1
    return day.replace(hour=end_h, minute=end_m, second=0)


def e(text):
    return html.escape(str(text or ""))


def ul(items):
    items = [i for i in items if i]
    return "<ul>" + "".join(f"<li>{e(i)}</li>" for i in items) + "</ul>" if items else "<p>-</p>"


def priority_badge(priority, original=None):
    cls = {"높음": "p-high", "보통": "p-mid", "낮음": "p-low"}.get(priority, "p-mid")
    if original and original != priority:
        order = ["낮음", "보통", "높음"]
        arrow = "↑" if order.index(priority) > order.index(original) else "↓"
        return f'<span class="badge {cls}">{e(priority)} {arrow} (요청: {e(original)})</span>'
    return f'<span class="badge {cls}">{e(priority)}</span>'


def build_item(t, cls):
    """티켓 1건의 보고서 표시용 데이터를 만든다. Router 경로별로 기본 문구가 다르다."""
    tid = t["ticket_id"]
    state, answer, review = ticket_state(tid, cls)
    route = cls["route"]
    path = WORK_DIR / "analysis" / f"{tid}.json"
    a = read_json(path) if path.exists() else {}
    rep = cls.get("representative_ticket")

    if route == "중복 VOC":
        a = {"evidence": [cls.get("reason", "")], "cause": f"선행 티켓({rep})과 같은 건",
             "actions": [f"{rep} 대표 티켓에 병합, 처리 결과는 대표 티켓으로 안내"],
             "action_owner": "직접 조치", "action_summary": f"{rep} 대표 티켓에 병합",
             "precondition": "", "notes": ["병합 처리. 별도 조치 없음."], **a}
    elif route == "개선 요청":
        a = {"evidence": ["해당 없음 (기능 추가 요청)"], "cause": "VOC 처리 대상이 아닌 기능 개선 요청",
             "actions": ["개선요청 게시판으로 이관하고 감사 인사와 이관 사실을 안내"],
             "action_owner": "이관", "transfer_to": "개선요청 게시판",
             "transfer_content": t["title"], "notes": ["개선요청 게시판 이관."], **a}

    received = parse_dt(t["received_at"])
    return {"t": t, "cls": cls, "a": a, "state": state, "answer": answer, "review": review,
            "deadline": deadline_of(received, cls["priority"])}


CSS = """
:root{--bg:#f5f6f8;--surface:#fff;--surface-2:#f0f2f5;--border:#dfe3e8;--text:#1d2330;--muted:#5d6675;--accent:#2457c5;
--high:#c62828;--high-bg:#fdecea;--mid:#a15c00;--mid-bg:#fff4e0;--low:#2e7d32;--low-bg:#e8f5e9;--info-bg:#e8effc;--warn-bg:#fff8e1;--warn:#8a6100}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#14171c;--surface:#1d2128;--surface-2:#262b33;--border:#353b45;--text:#e6e9ee;--muted:#9aa3b0;--accent:#7aa2f7;
--high:#ff8a80;--high-bg:#3a1f1f;--mid:#ffc46b;--mid-bg:#3a2e18;--low:#81c784;--low-bg:#1e3321;--info-bg:#1f2a40;--warn-bg:#3a3218;--warn:#ffd166}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font-family:"Pretendard","Malgun Gothic","Apple SD Gothic Neo",sans-serif;font-size:15px;line-height:1.6}
.wrap{max-width:1080px;margin:0 auto;padding:32px 16px 64px}header h1{margin:0 0 4px;font-size:26px}header p{margin:0;color:var(--muted)}
h2{font-size:18px;margin:36px 0 12px}.panel{background:var(--surface);border:1px solid var(--border);border-radius:12px;padding:18px 20px}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin-top:20px}
.stat{background:var(--surface);border:1px solid var(--border);border-radius:12px;padding:14px 16px}.stat .label{color:var(--muted);font-size:13px}
.stat .value{font-size:28px;font-weight:700}.stat .sub{font-size:13px;color:var(--muted)}
.two{display:grid;grid-template-columns:1fr 1fr;gap:12px}@media (max-width:760px){.two{grid-template-columns:1fr}}
.typebar{display:flex;flex-wrap:wrap;gap:6px}.chip{display:inline-block;padding:2px 10px;border-radius:999px;background:var(--surface-2);border:1px solid var(--border);font-size:13px}
table{width:100%;border-collapse:collapse;font-size:14px}th,td{text-align:left;padding:8px 6px;border-bottom:1px solid var(--border);vertical-align:top}
th{color:var(--muted);font-weight:600;font-size:13px}tr:last-child td{border-bottom:0}a{color:var(--accent);text-decoration:none}a:hover{text-decoration:underline}
.badge{display:inline-block;padding:1px 8px;border-radius:6px;font-size:12px;font-weight:700;white-space:nowrap}
.p-high{color:var(--high);background:var(--high-bg)}.p-mid{color:var(--mid);background:var(--mid-bg)}.p-low{color:var(--low);background:var(--low-bg)}
.b-type{color:var(--accent);background:var(--info-bg)}.b-flag{color:var(--warn);background:var(--warn-bg)}
.card{background:var(--surface);border:1px solid var(--border);border-radius:12px;margin-bottom:16px;overflow:hidden}
.card.failed{border:2px solid var(--high)}
.card-head{padding:14px 20px;border-bottom:1px solid var(--border);display:flex;flex-wrap:wrap;gap:8px 16px;align-items:baseline}
.card-head .no{font-weight:700;font-family:Consolas,monospace}.card-head .title{font-weight:700;font-size:16px;flex:1 1 280px}.card-head .meta{color:var(--muted);font-size:13px}
.card-body{padding:16px 20px;display:grid;grid-template-columns:1fr 1fr;gap:14px 24px}@media (max-width:760px){.card-body{grid-template-columns:1fr}}
.full{grid-column:1/-1}.sec h4{margin:0 0 4px;font-size:13px;color:var(--muted);font-weight:600}.sec ul{margin:0;padding-left:18px}.sec p{margin:0}
.voc{background:var(--surface-2);border-radius:8px;padding:10px 12px;font-size:14px;color:var(--muted)}
.note{border-left:3px solid var(--warn);background:var(--warn-bg);padding:8px 12px;border-radius:0 8px 8px 0}
.note.bad{border-left-color:var(--high);background:var(--high-bg)}
.answer{position:relative;border:1px solid var(--border);border-radius:8px;padding:12px 14px;background:var(--surface);white-space:pre-line}
.copy{position:absolute;top:8px;right:8px;font-size:12px;padding:3px 10px;border-radius:6px;border:1px solid var(--border);background:var(--surface-2);color:var(--text);cursor:pointer}
.copy:hover{border-color:var(--accent)}.answer-text{padding-right:56px}footer{margin-top:32px;color:var(--muted);font-size:13px}
"""

COPY_JS = """
document.querySelectorAll('.copy').forEach(function(btn){btn.addEventListener('click',function(){
var text=btn.parentElement.querySelector('.answer-text').innerText;
try{navigator.clipboard.writeText(text).then(function(){btn.textContent='복사됨';setTimeout(function(){btn.textContent='복사';},1500);},function(){});}catch(e){}});});
"""


def render_card(it):
    t, cls, a, ans, rev = it["t"], it["cls"], it["a"], it["answer"], it["review"]
    tid = t["ticket_id"]
    failed = it["state"] == "검토 미통과"
    flags = ""
    if cls["route"] == "중복 VOC" and cls.get("representative_ticket"):
        flags += f'<span class="badge b-flag">병합 → {e(cls["representative_ticket"])}</span>'
    if failed:
        flags += '<span class="badge p-high">검토 미통과</span>'
    elif it["state"] != "검토 통과":
        flags += f'<span class="badge b-flag">{e(it["state"])}</span>'

    if cls.get("category_changed"):
        type_text = f'요청자 카테고리 <b>{e(t["category"])}</b> → <b>{e(cls["type"])}</b>로 재분류. '
    else:
        type_text = f'카테고리 유지(<b>{e(t["category"])}</b>). '
    type_text += e(cls.get("priority_reason", ""))

    cause = f'<p>{e(a.get("cause"))}</p>' + ul(a.get("actions", []))
    notes = list(a.get("notes", []))
    if a.get("precondition"):
        notes.insert(0, f'선행 조건: {a["precondition"]}')
    if a.get("action_owner") == "이관" and a.get("transfer_to"):
        notes.insert(0, f'{a["transfer_to"]} 이관. 전달 내용: {a.get("transfer_content", "")}')
    notes_html = f'<div class="note">{"<br>".join(e(n) for n in notes) or "-"}</div>'

    if ans:
        answer_html = (f'<div class="answer"><button class="copy">복사</button>'
                       f'<div class="answer-text">{e(ans.get("answer"))}</div></div>')
    else:
        answer_html = '<div class="note bad">답변 초안이 없습니다.</div>'
    if failed and rev:
        issues = "".join(f'<li><b>{e(i.get("criterion"))}</b>: {e(i.get("detail"))}</li>' for i in rev.get("issues", []))
        answer_html += (f'<div class="note bad" style="margin-top:8px">재작성 {MAX_REVISIONS}회 후에도 검토를 통과하지 못했습니다. '
                        f'발송 전 직접 확인하세요.<ul>{issues}</ul></div>')
    elif ans:
        answer_html += f'<p style="margin-top:6px;font-size:13px;color:var(--muted)">답변 검토: {e(it["state"])} · 재작성 {ans.get("revision", 0)}회</p>'

    r = t["requester"]
    return f"""
<div class="card{' failed' if failed else ''}" id="{e(tid)}">
  <div class="card-head"><span class="no">{e(tid)}</span><span class="title">{e(t["title"])}</span>
    <span class="badge b-type">{e(cls["type"])}</span>{priority_badge(cls["priority"], t["priority"])}{flags}
    <span class="meta">{e(r.get("name"))} · {e(r.get("dept"))} · {e(t["received_at"][11:16])} 접수 · 기한 {it["deadline"]:%m-%d %H:%M}</span></div>
  <div class="card-body">
    <div class="sec full"><h4>VOC 내용</h4><div class="voc">{e(t["content"])}</div></div>
    <div class="sec"><h4>처리 유형·우선순위</h4><p>{type_text}</p></div>
    <div class="sec"><h4>확인 근거 (DB/로그)</h4>{ul(a.get("evidence", []))}</div>
    <div class="sec"><h4>원인·조치 방안</h4>{cause}</div>
    <div class="sec"><h4>이관·주의사항</h4>{notes_html}</div>
    <div class="sec full"><h4>답변 초안</h4>{answer_html}</div>
  </div>
</div>"""


def cmd_report():
    tickets = read_json(WORK_DIR / "tickets.json")["tickets"]
    classified = load_classified()
    items = [build_item(t, classified[t["ticket_id"]]) for t in tickets]
    now = datetime.now()

    pending = [it for it in items if it["state"] not in ("검토 통과", "검토 미통과")]
    if pending:
        print("아직 끝나지 않은 티켓이 있습니다 (보고서에는 현재 상태로 표시):")
        for it in pending:
            print(f"  {it['t']['ticket_id']}: {it['state']}")

    def count(key):
        out = {}
        for it in items:
            out[key(it)] = out.get(key(it), 0) + 1
        return out

    by_type = count(lambda it: it["cls"]["type"])
    by_pri = count(lambda it: it["cls"]["priority"])
    direct = [it for it in items if it["a"].get("action_owner") == "직접 조치"]
    transfer = [it for it in items if it["a"].get("action_owner") == "이관"]
    changed = sum(1 for it in items if it["cls"].get("category_changed"))
    failed = [it for it in items if it["state"] == "검토 미통과"]
    received = sorted(parse_dt(it["t"]["received_at"]) for it in items)

    def link(it):
        tid = it["t"]["ticket_id"]
        return f'<a href="#{e(tid)}">{e(tid)}</a>'

    def deadline_cell(it):
        d = it["deadline"]
        tag = ""
        if now > d:
            tag = ' <span class="badge p-high">기한 초과</span>'
        elif d - now <= timedelta(minutes=SOON_MINUTES):
            tag = ' <span class="badge p-high">임박</span>'
        return f"{d:%m-%d %H:%M}까지{tag}"

    first = sorted([it for it in items if it["cls"]["priority"] == "높음"], key=lambda it: it["deadline"])
    first_rows = "".join(
        f'<tr><td>{link(it)}</td><td>{e(it["cls"].get("summary") or it["t"]["title"])}</td>'
        f'<td>{priority_badge(it["cls"]["priority"], it["t"]["priority"])}</td><td>{deadline_cell(it)}</td></tr>'
        for it in first) or '<tr><td colspan="4">우선순위 높음 건이 없습니다.</td></tr>'
    direct_rows = "".join(
        f'<tr><td>{link(it)}</td><td>{e(it["a"].get("action_summary") or "; ".join(it["a"].get("actions", [])))}</td>'
        f'<td>{e(it["a"].get("precondition") or "-")}</td></tr>' for it in direct) or '<tr><td colspan="3">-</td></tr>'
    transfer_rows = "".join(
        f'<tr><td>{link(it)}</td><td>{e(it["a"].get("transfer_to"))}</td><td>{e(it["a"].get("transfer_content"))}</td></tr>'
        for it in transfer) or '<tr><td colspan="3">-</td></tr>'
    order = {"높음": 0, "보통": 1, "낮음": 2}
    cards = "".join(render_card(it) for it in sorted(items, key=lambda it: (order.get(it["cls"]["priority"], 1), it["deadline"])))
    type_chips = "".join(f'<span class="chip">{e(k)} {v}</span>' for k, v in by_type.items())
    pri_chips = " ".join(f'<span class="badge {c}">{p} {by_pri.get(p, 0)}</span>'
                         for p, c in (("높음", "p-high"), ("보통", "p-mid"), ("낮음", "p-low")))
    failed_stat = (f'<div class="stat"><div class="label">검토 미통과</div><div class="value" style="color:var(--high)">{len(failed)}건</div>'
                   f'<div class="sub">발송 전 직접 확인</div></div>') if failed else ""

    page = f"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>VOC 처리 검토 보고서</title><style>{CSS}</style></head>
<body><div class="wrap">
<header><h1>VOC 처리 검토 보고서</h1>
<p>{now:%Y-%m-%d %H:%M} 기준 · 접수 VOC {len(items)}건 · 작성: VOC 처리 Agent (검토 대기)</p></header>
<div class="stats">
  <div class="stat"><div class="label">오늘 접수</div><div class="value">{len(items)}건</div><div class="sub">{received[0]:%H:%M} ~ {received[-1]:%H:%M} 접수</div></div>
  <div class="stat"><div class="label">우선순위 높음</div><div class="value" style="color:var(--high)">{by_pri.get("높음", 0)}건</div><div class="sub">보통 {by_pri.get("보통", 0)} · 낮음 {by_pri.get("낮음", 0)}</div></div>
  <div class="stat"><div class="label">직접 조치</div><div class="value">{len(direct)}건</div><div class="sub">확인 후 처리 가능</div></div>
  <div class="stat"><div class="label">이관 필요</div><div class="value">{len(transfer)}건</div><div class="sub">{e(" · ".join(sorted({it["a"].get("transfer_to", "") for it in transfer})) or "-")}</div></div>
  <div class="stat"><div class="label">카테고리 재분류</div><div class="value">{changed}건</div><div class="sub">요청자 선택과 다름</div></div>
  {failed_stat}
</div>
<h2>오늘 VOC 현황</h2>
<div class="two">
  <div class="panel"><h4 style="margin:0 0 10px">처리 유형별</h4><div class="typebar">{type_chips}</div></div>
  <div class="panel"><h4 style="margin:0 0 10px">우선순위별</h4><div class="typebar">{pri_chips}</div>
  <p style="margin:10px 0 0;font-size:13px;color:var(--muted)">처리 기한: 높음 {SLA_HIGH_HOURS}시간 이내 · 보통 당일 · 낮음 {SLA_LOW_BUSINESS_DAYS}영업일 이내</p></div>
</div>
<h2>먼저 볼 건</h2>
<div class="panel"><table><tr><th style="width:130px">티켓</th><th>내용</th><th style="width:150px">우선순위</th><th style="width:170px">처리 기한</th></tr>{first_rows}</table></div>
<h2>내가 할 조치 / 이관 목록</h2>
<div class="two">
  <div class="panel"><h4 style="margin:0 0 6px">직접 조치</h4><table><tr><th style="width:120px">티켓</th><th>조치</th><th style="width:110px">선행 조건</th></tr>{direct_rows}</table></div>
  <div class="panel"><h4 style="margin:0 0 6px">이관</h4><table><tr><th style="width:120px">티켓</th><th>이관처</th><th>전달 내용</th></tr>{transfer_rows}</table></div>
</div>
<h2>티켓별 검토</h2>
{cards}
<footer>답변 초안은 담당자 검토 후 등록합니다. 자동 발송되지 않습니다.</footer>
</div><script>{COPY_JS}</script></body></html>
"""
    out = OUTPUT_DIR / f"VOC_검토보고서_{now:%Y%m%d}.html"
    out.write_text(page, encoding="utf-8")
    print(f"보고서 → {out.relative_to(BASE_DIR)}")


COMMANDS = {"load": cmd_load, "context": cmd_context, "status": cmd_status, "report": cmd_report}

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 2 or sys.argv[1] not in COMMANDS:
        sys.exit("사용법: python scripts/voc_tools.py [load|context|status|report]")
    COMMANDS[sys.argv[1]]()
