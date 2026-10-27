"""검사로그 + 설비센서로그를 LOT_ID로 결합해 이상데이터를 엑셀로 정리"""
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── 이상 기준 ──
DEFECT_RATE_LIMIT = 3.0      # 1) 불량률(%) >= 3.0
EXCURSION_RATIO = 0.6        # 3) 최대/최소가 정상범위를 (상한-하한)×0.6 넘게 벗어남

# 사용법: python find_anomaly.py [데이터폴더] [출력파일]  (기본: data → 이상데이터_분석.xlsx)
data_dir = sys.argv[1] if len(sys.argv) > 1 else "data"
out_path = sys.argv[2] if len(sys.argv) > 2 else "이상데이터_분석.xlsx"
insp = pd.read_csv(f"{data_dir}/검사로그.csv", encoding="utf-8-sig")
sens = pd.read_csv(f"{data_dir}/설비센서로그.csv", encoding="utf-8-sig")

df = insp.merge(sens.drop(columns=["설비ID", "라인", "공정"]), on="LOT_ID", how="outer", validate="1:1")
rng = df["정상범위"].str.split("~", expand=True).astype(float)
df["하한"], df["상한"] = rng[0], rng[1]
margin = (df["상한"] - df["하한"]) * EXCURSION_RATIO

c1 = df["불량률(%)"] >= DEFECT_RATE_LIMIT
c2 = (df["평균값"] < df["하한"]) | (df["평균값"] > df["상한"])
c3_max = df["최대값"] > df["상한"] + margin
c3_min = df["최소값"] < df["하한"] - margin
c3 = c3_max | c3_min

def reasons(i):
    r = []
    if c1[i]: r.append(f"불량률 {df.at[i,'불량률(%)']}% ≥ {DEFECT_RATE_LIMIT}")
    if c2[i]: r.append(f"평균값 {df.at[i,'평균값']} 정상범위({df.at[i,'정상범위']}) 이탈")
    if c3_max[i]: r.append(f"최대값 {df.at[i,'최대값']} > 상한+{margin[i]:.3g}")
    if c3_min[i]: r.append(f"최소값 {df.at[i,'최소값']} < 하한-{margin[i]:.3g}")
    return " / ".join(r)

df["①불량률"] = c1.map({True: "●", False: ""})
df["②평균이탈"] = c2.map({True: "●", False: ""})
df["③극값이탈"] = c3.map({True: "●", False: ""})
df["해당기준수"] = c1.astype(int) + c2.astype(int) + c3.astype(int)
anom = df[c1 | c2 | c3].copy()
anom["이상사유"] = [reasons(i) for i in anom.index]
anom = anom.sort_values(["검사일자", "LOT_ID"])

cols = ["검사일자", "라인", "설비ID", "공정", "LOT_ID", "검사수량", "불량수량", "불량률(%)", "불량유형", "검사자",
        "모니터링항목", "평균값", "최대값", "최소값", "단위", "정상범위", "가동시간(hr)", "가동율(%)",
        "①불량률", "②평균이탈", "③극값이탈", "해당기준수", "이상사유"]
anom = anom[cols]

# 요약
n = len(df)
summary = pd.DataFrame([
    ["① 불량률(%) ≥ 3.0", int(c1.sum())],
    ["② 평균값이 정상범위(하한~상한) 벗어남", int(c2.sum())],
    ["③ 최대/최소값이 정상범위를 (상한-하한)×0.6 넘게 벗어남", int(c3.sum())],
    ["이상 LOT 합계 (하나라도 해당, 중복 제외)", len(anom)],
    ["전체 LOT 수", n],
    ["이상 비율(%)", round(len(anom) / n * 100, 2)],
], columns=["구분", "건수"])

def group(key):
    g = df.groupby(key).agg(전체LOT=("LOT_ID", "size"),
                            불량률이상=("①불량률", lambda s: (s == "●").sum()),
                            평균이탈=("②평균이탈", lambda s: (s == "●").sum()),
                            극값이탈=("③극값이탈", lambda s: (s == "●").sum()),
                            이상LOT=("해당기준수", lambda s: (s > 0).sum()))
    g["이상비율(%)"] = (g["이상LOT"] / g["전체LOT"] * 100).round(2)
    return g.sort_values("이상LOT", ascending=False).reset_index()

by_eq = group(["설비ID", "공정", "라인"])
by_item = group(["공정", "모니터링항목", "정상범위"])
df["연월"] = df["검사일자"].str[:7]
by_month = group("연월").sort_values("연월")

with pd.ExcelWriter(out_path, engine="openpyxl") as w:
    summary.to_excel(w, sheet_name="요약", index=False)
    anom.to_excel(w, sheet_name="이상데이터", index=False)
    by_eq.to_excel(w, sheet_name="설비별", index=False)
    by_item.to_excel(w, sheet_name="모니터링항목별", index=False)
    by_month.to_excel(w, sheet_name="월별추이", index=False)

# 서식
wb = load_workbook(out_path)
hdr_fill = PatternFill("solid", fgColor="1F4E78")
red = PatternFill("solid", fgColor="F8CBAD")
for ws in wb.worksheets:
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value or "")) for c in list(col)[:300])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width * 1.6 + 2, 8), 70)
    if ws.title != "요약":
        ws.auto_filter.ref = ws.dimensions

ws = wb["이상데이터"]
h = {c.value: c.column for c in ws[1]}
for r, (_, row) in enumerate(anom.iterrows(), start=2):
    lo, hi = map(float, row["정상범위"].split("~"))
    m = (hi - lo) * EXCURSION_RATIO
    if row["불량률(%)"] >= DEFECT_RATE_LIMIT: ws.cell(r, h["불량률(%)"]).fill = red
    if not lo <= row["평균값"] <= hi: ws.cell(r, h["평균값"]).fill = red
    if row["최대값"] > hi + m: ws.cell(r, h["최대값"]).fill = red
    if row["최소값"] < lo - m: ws.cell(r, h["최소값"]).fill = red
    for k in ("①불량률", "②평균이탈", "③극값이탈", "해당기준수"):
        ws.cell(r, h[k]).alignment = Alignment(horizontal="center")

ws = wb["요약"]
ws.column_dimensions["A"].width = 60
ws.cell(len(summary) + 3, 1, f"기준값: 불량률 ≥ {DEFECT_RATE_LIMIT}%, 극값 허용폭 = (상한-하한) × {EXCURSION_RATIO}").font = Font(italic=True, color="666666")
ws.cell(len(summary) + 4, 1, "원천: data/검사로그.csv + data/설비센서로그.csv (LOT_ID 1:1 결합)").font = Font(italic=True, color="666666")
wb.save(out_path)

print(summary.to_string(index=False))
print("\n[설비별 상위]\n", by_eq.head(8).to_string(index=False))
print("\n[항목별]\n", by_item.to_string(index=False))
